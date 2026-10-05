"""Create readable DOCX and PDF versions of a prepared text draft."""

import re
from pathlib import Path
from xml.sax.saxutils import escape


def export_document_versions(text_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text_path.read_text(encoding="utf-8"))
    title = "Resume draft" if text_path.stem.startswith("resume_") else "Cover letter"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    doc.core_properties.author = ""
    doc.core_properties.title = title
    doc.add_heading(title, level=0)

    font = "Helvetica"
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    if font_path.exists():
        font = "JobHunterArial"
        if font not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(font, str(font_path)))
    body_style = ParagraphStyle("DraftBody", fontName=font, fontSize=10.5, leading=14,
                                spaceAfter=5, textColor=colors.HexColor("#172331"))
    heading_style = ParagraphStyle("DraftHeading", parent=body_style, fontSize=12,
                                  spaceBefore=10, spaceAfter=6, keepWithNext=True)
    title_style = ParagraphStyle("DraftTitle", parent=body_style, fontSize=18, leading=22, spaceAfter=14)
    flow = [Paragraph(title, title_style)]
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            flow.append(Spacer(1, 4))
            continue
        heading = line.endswith(":") or line.startswith("--- Original resume")
        if line.startswith("--- Original resume"):
            line = "Experience and background"
        if heading:
            doc.add_heading(line.rstrip(":"), level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line)
        flow.append(Paragraph(escape(line), heading_style if heading else body_style))
    doc.save(str(text_path.with_suffix(".docx")))
    SimpleDocTemplate(str(text_path.with_suffix(".pdf")), pagesize=A4,
                      rightMargin=50, leftMargin=50, topMargin=46, bottomMargin=46,
                      title=title, author="").build(flow)
